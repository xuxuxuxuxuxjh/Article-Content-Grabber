from docx import Document
import xlwt
import os

wb = xlwt.Workbook()
st = wb.add_sheet('sheet')
# items = ['答题人', '学号', '课程', '班级', '提交时间', 'ip', '考试得分']
# multi_items = ['学生答案']

once_items = list(input().split(' '))
multi_items = list(input().split(' '))
once_ans = [0]*len(once_items)
multi_ans = [0]*len(multi_items)
present = ["false"]*len(multi_items)
# print(once_items)
# print(multi_items)

def init():
    for i in range(len(once_items)):
        st.col(i).width = 256*15   
        st.write(0, i, once_items[i])

    id = len(once_items)
    for i in range(0,len(multi_items),2):
        st.col(id).width = 256*15
        st.write(0, id, multi_items[i]+' '+multi_items[i+1])
        id = id + 1;


def grab(name, row, postfix):
    for i in range(len(once_ans)):
        once_ans[i]=0
    for i in range(len(multi_ans)):
        multi_ans[i]=0
    for i in range(len(present)):
        present[i]='false'

    docx = Document(name+postfix)
    pg = [];
    for i in docx.paragraphs:
        pg.append(i.text)

    for text in pg:
        list_first = list(text.split('：'))
        list_last = []
        for first in list_first:
            list_last.extend(first.split(' '))

        data=[]
        for i in range(len(list_last)):
            list_last[i].strip(' ')
            if list_last[i] != '':
                data.append(list_last[i])
                
        for i in range(len(data)):
            for j in range(len(once_items)):
                if(data[i] == once_items[j]):
                    once_ans[j] = data[i+1]
        
        for i in range(len(data)):
            for j in range(0,len(multi_items),2):
                if data[i].find(multi_items[j]):
                    present[j] = "true"
            for j in range(0,len(multi_items),2):
                if data[i] == multi_items[j+1] and present[j] == "true":
                    present[j] = "false"
                    multi_ans[j] = '：'.join(data[i+1:])
    for i in range(len(once_ans)):
        st.write(row, i, once_ans[i])
    id = len(once_ans)
    for i in range(0,len(multi_ans),2):
        st.write(row, id, multi_ans[i])
        id = id + 1


init()
path = './'
files = os.listdir(path)
row = 0
for file in files:
    lst = list(file.split('.'))
    if lst[-1] != 'docx' :
        continue
    row = row + 1
    name = '.'.join(lst[:-1])
    print(row)
    grab(name, row, '.'+lst[-1])

wb.save('Sheet.xls')


#答题人 学号 课程 班级 提交时间 ip 考试得分
#简答题 学生答案
