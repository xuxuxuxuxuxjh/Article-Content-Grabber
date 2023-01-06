import re
from docx import Document
import xlwt
import os
from win32com import client as wc

wb = xlwt.Workbook()
st = wb.add_sheet('sheet1')
items = ['答题人', '学号', '课程', '班级', '提交时间', 'ip', '考试得分']
spj = '学生答案'

def init():
    for i in range(len(items)):
        st.col(i).width = 256*15
        st.write(0, i, items[i])
    st.col(len(items)).width = 256*15
    st.write(0, len(items), spj)


def save_as_docx(file):
    word = wc.Dispatch('Word.Application')
    doc = word.Documents.Open(file)        # 目标路径下的文件
    doc.SaveAs(file+'x', 12, False, "", True, "", False, False, False, False)  # 转化后路径下的文件    
    doc.Close()
    word.Quit()

def add_file(name, row, suffix):
    if suffix == '.doc':
        save_as_docx(os.getcwd()+'\\'+name+'.doc')
        suffix = '.docx'
    docx = Document(name+suffix)
    pg = [paragraph.text for paragraph in docx.paragraphs]
    data = []

    flag = False
    for text in pg:
        first = text.split('：')
        lst = []
        for list_first in first:
            # print(type(list_first.split(' ')))
            lst.extend(list_first.split(' ')) 
        # print(lst)
        # lst = re.split('[： ]', text)
        dat = []
        if text.find('简答题') != -1:
            flag = True
        for i in range(len(lst)):
            lst[i].strip(' ')
            if lst[i] != '':
                dat.append(lst[i])

        for i in range(len(dat)):
            if dat[i] in items:
                data.append(dat[i+1])
            if flag and dat[i] == spj:
                data.append('：'.join(dat[i+1:]))
                flag = False

    # print(row)
    for i in range(len(data)):
        st.write(row, i, data[i])

# print(os.getcwd())
init()
path = './'
files = os.listdir(path)
row = 0
for file in files:
    lst = list(file.split('.'))
    if lst[-1] != 'docx' and lst[-1] != 'doc':
        continue
    row += 1
    name = '.'.join(lst[:-1])
    # print(name)
    print(row)
    add_file(name, row, '.'+lst[-1])

wb.save('Sheet.xls')

